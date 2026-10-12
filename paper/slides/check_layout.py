# -*- coding: utf-8 -*-
"""Static layout checker: flag out-of-canvas shapes and text likely to overflow.

Heuristic only, but it catches the two real defect classes seen in rendering:
text spilling past its box, and elements crossing the slide edge.
"""

import io
import math
import os
import sys
import unicodedata

from pptx import Presentation
from pptx.util import Emu

sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding="utf-8",
                              errors="replace", line_buffering=True)

HERE = os.path.dirname(os.path.abspath(__file__))
PPTX = os.path.join(HERE, "output", "identifiability-boundary-popular.pptx")
W_IN, H_IN = 10.0, 5.625
EMU_IN = 914400.0


def is_wide(ch):
    return unicodedata.east_asian_width(ch) in ("W", "F")


def text_width_units(txt):
    """Width in em units: CJK/full-width = 1.0, others = 0.52."""
    u = 0.0
    for ch in txt:
        u += 1.0 if is_wide(ch) else 0.52
    return u


def para_info(p):
    runs = p.runs
    if not runs:
        return None
    txt = "".join(r.text for r in runs)
    if not txt.strip():
        return None
    size = None
    for r in runs:
        if r.font.size is not None:
            size = max(size or 0, r.font.size.pt)
    if size is None:
        size = 18.0
    spacing = p.line_spacing if isinstance(p.line_spacing, float) else 1.0
    before = p.space_before.pt if p.space_before is not None else 0.0
    after = p.space_after.pt if p.space_after is not None else 0.0
    return txt, size, spacing, before, after


def check():
    prs = Presentation(PPTX)
    problems = []
    for idx, slide in enumerate(prs.slides, start=1):
        for sh in slide.shapes:
            x = sh.left / EMU_IN
            y = sh.top / EMU_IN
            w = sh.width / EMU_IN
            h = sh.height / EMU_IN
            name = sh.shape_type

            if x < -0.001 or y < -0.001 or x + w > W_IN + 0.001 or y + h > H_IN + 0.001:
                problems.append(
                    (idx, "CANVAS", "shape out of bounds",
                     "x=%.2f y=%.2f w=%.2f h=%.2f" % (x, y, w, h)))

            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            ml = (tf.margin_left or 0) / EMU_IN
            mr = (tf.margin_right or 0) / EMU_IN
            usable_w = max(w - ml - mr, 0.1)

            total_pt = 0.0
            for p in tf.paragraphs:
                info = para_info(p)
                if info is None:
                    continue
                txt, size, spacing, before, after = info
                units = text_width_units(txt)
                text_w_in = units * size / 72.0
                lines = max(1, math.ceil(text_w_in / usable_w - 1e-6))
                line_h_pt = size * 1.22 * spacing
                total_pt += before + after + lines * line_h_pt

            need_in = total_pt / 72.0
            if need_in > h + 0.02:
                problems.append(
                    (idx, "OVERFLOW", "text taller than its box",
                     "need=%.2f\" box=%.2f\"  box(x=%.2f y=%.2f w=%.2f h=%.2f)  text=%r"
                     % (need_in, h, x, y, w, h, tf.text[:46])))
    return problems


if __name__ == "__main__":
    probs = check()
    if not probs:
        print("LAYOUT OK - no canvas overflow or text overflow detected")
    else:
        cur = None
        for idx, kind, msg, detail in probs:
            if idx != cur:
                print("\n--- slide %02d ---" % idx)
                cur = idx
            print("  [%s] %s" % (kind, detail))
        print("\ntotal problems:", len(probs))
    sys.exit(0)
