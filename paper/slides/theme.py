# -*- coding: utf-8 -*-
"""Design tokens and layout helpers shared by every slide module.

Canvas is LAYOUT_16x9 -> 10.0 in x 5.625 in.
Palette is the "Education & Charts" set from the deck design system.
"""

import os

from lxml import etree
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt

W = 10.0
H = 5.625
IN = Inches

FONT = "Microsoft YaHei"
FONT_NUM = "Georgia"

PALETTE = {
    "primary": "006D77",   # deep teal
    "secondary": "83C5BE",  # soft teal
    "accent": "E29578",   # terracotta
    "light": "FFDDD2",   # pale peach
    "bg": "EDF6F9",   # pale blue white
}

ALIGN = {
    "left": PP_ALIGN.LEFT,
    "center": PP_ALIGN.CENTER,
    "right": PP_ALIGN.RIGHT,
    "justify": PP_ALIGN.JUSTIFY,
}
ANCHOR = {
    "top": MSO_ANCHOR.TOP,
    "middle": MSO_ANCHOR.MIDDLE,
    "bottom": MSO_ANCHOR.BOTTOM,
}
KIND = {
    "rect": MSO_SHAPE.RECTANGLE,
    "round": MSO_SHAPE.ROUNDED_RECTANGLE,
    "oval": MSO_SHAPE.OVAL,
}

BASE = os.path.dirname(os.path.abspath(__file__))
FIGDIR = os.path.abspath(os.path.join(BASE, "..", "figures"))
ANIMEDIR = os.path.join(FIGDIR, "anime")
ASSETS = os.path.join(BASE, "assets")


def anime(name):
    return os.path.join(ANIMEDIR, name)


def asset(name):
    return os.path.join(ASSETS, name)


def rgb(key):
    return RGBColor.from_string(PALETTE.get(key, key))


def new_slide(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def set_bg(slide, color="bg"):
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(color)


def shape(slide, x, y, w, h, fill=None, line=None, lw=1.0, kind="rect", radius=None):
    s = slide.shapes.add_shape(KIND[kind], Inches(x), Inches(y), Inches(w), Inches(h))
    if radius is not None and kind == "round":
        try:
            s.adjustments[0] = radius
        except Exception:
            pass
    if fill:
        s.fill.solid()
        s.fill.fore_color.rgb = rgb(fill)
    else:
        s.fill.background()
    if line:
        s.line.color.rgb = rgb(line)
        s.line.width = Pt(lw)
    else:
        s.line.fill.background()
    try:
        s.shadow.inherit = False
    except Exception:
        pass
    return s


def tbox(slide, x, y, w, h, anchor="top", margin=0.0):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = ANCHOR[anchor]
    tf.margin_left = Inches(margin)
    tf.margin_right = Inches(margin)
    tf.margin_top = Inches(0)
    tf.margin_bottom = Inches(0)
    return tf


def _ea(run, font):
    rPr = run._r.get_or_add_rPr()
    for tag in ("a:latin", "a:ea", "a:cs"):
        el = rPr.find(qn(tag))
        if el is None:
            el = etree.SubElement(rPr, qn(tag))
        el.set("typeface", font)


def para(tf, text, size=13, color="primary", bold=False, align="left", font=FONT,
         before=0, after=0, line=None, first=False, italic=False):
    p = tf.paragraphs[0] if first else tf.add_paragraph()
    p.alignment = ALIGN[align]
    p.space_before = Pt(before)
    p.space_after = Pt(after)
    if line:
        p.line_spacing = line
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.italic = italic
    r.font.color.rgb = rgb(color)
    r.font.name = font
    _ea(r, font)
    return p


def dot(slide, cx, cy, d=0.10, color="accent"):
    return shape(slide, cx - d / 2.0, cy - d / 2.0, d, d, fill=color, kind="oval")


def badge(slide, n, dark=False):
    tf = tbox(slide, 9.10, 5.04, 0.70, 0.30, anchor="middle")
    para(tf, "%02d" % n, size=11, color=("light" if dark else "primary"),
         align="right", font=FONT_NUM, first=True)


def header(slide, eyebrow, title, n):
    """Standard content-slide frame: left accent bar, eyebrow, title, corner mark."""
    set_bg(slide, "bg")
    shape(slide, 0, 0, 0.16, H, fill="accent")
    dot(slide, 0.60, 0.42, d=0.11, color="accent")
    tf = tbox(slide, 0.78, 0.30, 7.4, 0.26)
    para(tf, eyebrow, size=10.5, color="primary", bold=True, first=True)
    tf = tbox(slide, 0.55, 0.58, 7.6, 0.52)
    para(tf, title, size=25, color="primary", bold=True, first=True)
    shape(slide, 8.80, 0.20, 0.92, 0.92, line="secondary", lw=1.25, kind="oval")
    shape(slide, 9.13, 0.53, 0.24, 0.24, fill="accent", kind="oval")
    badge(slide, n)


def card(slide, x, y, w, h, style="ghost", lw=1.0):
    if style == "peach":
        return shape(slide, x, y, w, h, fill="light", kind="round", radius=0.10)
    if style == "teal":
        return shape(slide, x, y, w, h, fill="primary", kind="round", radius=0.10)
    return shape(slide, x, y, w, h, fill=None, line="secondary", lw=lw,
                 kind="round", radius=0.10)


def source(slide, text, y=5.04, x=0.55, w=6.4):
    tf = tbox(slide, x, y, w, 0.28)
    para(tf, text, size=10, color="primary", first=True)
